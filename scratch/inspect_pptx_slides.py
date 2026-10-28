import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

prs = Presentation('presentation/final_project_presentation.pptx')

for s_num in [1]:
    slide = prs.slides[s_num - 1]
    print(f'============================== SLIDE {s_num} ==============================')
    for shape in slide.shapes:
        if shape.has_text_frame:
            print(f'Shape {shape.shape_id} ({shape.name}):')
            for p_idx, p in enumerate(shape.text_frame.paragraphs):
                print(f'  Paragraph {p_idx}: "{p.text}"')
                for r_idx, r in enumerate(p.runs):
                    font_name = r.font.name
                    font_size = r.font.size.pt if r.font.size else None
                    bold = r.font.bold
                    color = str(r.font.color.rgb) if r.font.color and hasattr(r.font.color, 'rgb') and r.font.color.type is not None else None
                    print(f'    Run {r_idx}: "{r.text}" [font={font_name}, size={font_size}, bold={bold}, color={color}]')
