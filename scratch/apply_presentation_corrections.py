"""
apply_presentation_corrections.py
=================================
Applies the 4 targeted corrections to the PowerPoint presentation:
1. Slide 1: Correct author details (Sai Charan, 24JE1093, Applied Remote Sensing Practical, Applied Geophysics IIT (ISM) Dhanbad).
2. Slide 2: Replace 82.0% with 94.4% in curated dataset statistic.
3. Slide 8: Replace M3 explanation with qualified scientific text.
4. Slide 10: Replace 82% with 94.4% in curated dataset and replace M3 conclusion with qualified text.

Saves output as:
- C:\\Users\\manda\\Downloads\\Remote Sensing\\final_project_presentation_corrected.pptx
- and updates existing presentation files.
"""

import os
import sys
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

def apply_corrections():
    src_path = os.path.abspath("presentation/final_project_presentation.pptx")
    if not os.path.exists(src_path):
        src_path = os.path.abspath("final_project_presentation.pptx")
    
    prs = Presentation(src_path)
    print(f"Loaded presentation from: {src_path} ({len(prs.slides)} slides)")

    # -------------------------------------------------------------------------
    # Correction 1: Slide 1 Author Details
    # -------------------------------------------------------------------------
    slide1 = prs.slides[0]
    tb1 = [s for s in slide1.shapes if s.shape_id == 4][0]  # TextBox 3
    p2 = tb1.text_frame.paragraphs[2]
    
    # Clear runs and rebuild with exact required info
    p2.text = ""  # clears all runs in paragraph
    
    r1 = p2.add_run()
    r1.text = "Author: Sai Charan (Mandala Sai Charan)  |  Admission No.: 24JE1093\n"
    r1.font.size = Pt(13)
    r1.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
    
    r2 = p2.add_run()
    r2.text = "Course: Applied Remote Sensing Practical  |  Department: Applied Geophysics, IIT (ISM) Dhanbad\n"
    r2.font.size = Pt(13)
    r2.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
    
    r3 = p2.add_run()
    r3.text = "Dataset: BigEarthNet v2.0 (Sentinel-2 Level-2A BOA Surface Reflectance)  |  Date: October 2026"
    r3.font.size = Pt(13)
    r3.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    print("[OK] Slide 1: Author details updated.")

    # -------------------------------------------------------------------------
    # Correction 2: Slide 2 Dataset Statistic (82% -> 94.4%)
    # -------------------------------------------------------------------------
    slide2 = prs.slides[1]
    tb2 = [s for s in slide2.shapes if s.shape_id == 5][0]  # TextBox 4 (The Multi-Label Paradigm)
    p2_s2 = tb2.text_frame.paragraphs[2]
    p2_s2.text = ""
    
    r2_0 = p2_s2.add_run()
    r2_0.text = "• 94.4% Empirical Multi-Label Prevalence:"
    r2_0.font.bold = True
    r2_0.font.size = Pt(13)
    r2_0.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    
    r2_1 = p2_s2.add_run()
    r2_1.text = " 94.4% of patches in our curated dataset contain at least two land-cover classes."
    r2_1.font.bold = False
    r2_1.font.size = Pt(13)
    r2_1.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    print("[OK] Slide 2: Multi-label prevalence updated to 94.4% for curated dataset.")

    # -------------------------------------------------------------------------
    # Correction 3: Slide 8 M3 Performance Explanation
    # -------------------------------------------------------------------------
    slide8 = prs.slides[7]
    tb8 = [s for s in slide8.shapes if s.shape_id == 5][0]  # TextBox 4
    p3_s8 = tb8.text_frame.paragraphs[3]
    p3_s8.text = ""
    
    r8_0 = p3_s8.add_run()
    r8_0.text = "• M3 (Multispectral):"
    r8_0.font.bold = True
    r8_0.font.size = Pt(10.5)
    r8_0.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    
    r8_1 = p3_s8.add_run()
    r8_1.text = " Macro-F1 = 0.5673, Micro-F1 = 0.7081, Exact Match = 5.0%. The six-band model (M3) underperformed compared with the RGB+NIR model (M2) on the observed test metrics. Possible explanations include the small training sample, spectral collinearity, resampling effects, and optimization instability. These explanations were not independently tested."
    r8_1.font.bold = False
    r8_1.font.size = Pt(10.5)
    r8_1.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    print("[OK] Slide 8: M3 explanation qualified with exact scientific phrasing.")

    # -------------------------------------------------------------------------
    # Correction 4: Slide 10 Statistic & Cautious M3 Conclusion
    # -------------------------------------------------------------------------
    slide10 = prs.slides[9]
    tb10 = [s for s in slide10.shapes if s.shape_id == 5][0]  # TextBox 4
    p1_s10 = tb10.text_frame.paragraphs[1]
    p1_s10.text = ""
    
    r10_0 = p1_s10.add_run()
    r10_0.text = "• Multi-label recognition is vital: 94.4% of patches in our curated dataset contain at least two land-cover classes.\n\n"
    r10_0.font.size = Pt(10.5)
    r10_0.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    
    r10_1 = p1_s10.add_run()
    r10_1.text = "• NIR is the most valuable spectral addition: M2 (Visible+NIR) gained +11.32 pp Macro-F1 and tripled Exact Match (38%).\n\n"
    r10_1.font.size = Pt(10.5)
    r10_1.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    
    r10_2 = p1_s10.add_run()
    r10_2.text = "• Performance of M3: The six-band model (M3) performed worse than the RGB+NIR model (M2) on the observed test metrics. Small sample size, spectral collinearity, resampling effects, and optimization instability are possible explanations, but they were not independently tested."
    r10_2.font.size = Pt(10.5)
    r10_2.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    print("[OK] Slide 10: Conclusion updated with 94.4% and cautious M3 interpretation.")

    # -------------------------------------------------------------------------
    # Save outputs
    # -------------------------------------------------------------------------
    out_corrected = os.path.abspath("final_project_presentation_corrected.pptx")
    prs.save(out_corrected)
    print(f"Saved: {out_corrected} ({os.path.getsize(out_corrected):,} bytes)")
    
    out_root = os.path.abspath("final_project_presentation.pptx")
    prs.save(out_root)
    print(f"Saved: {out_root} ({os.path.getsize(out_root):,} bytes)")

    out_sub = os.path.abspath("presentation/final_project_presentation.pptx")
    try:
        prs.save(out_sub)
        print(f"Saved: {out_sub} ({os.path.getsize(out_sub):,} bytes)")
    except Exception as e:
        print(f"Note on {out_sub}: {e}")

    return out_corrected

if __name__ == "__main__":
    apply_corrections()
