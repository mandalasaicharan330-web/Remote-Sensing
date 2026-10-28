import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

prs = Presentation('presentation/final_project_presentation.pptx')
slide = prs.slides[0]
for idx, shape in enumerate(slide.shapes):
    print(f'Shape index {idx}, ID {shape.shape_id}, Name: {shape.name}')
    if shape.has_text_frame:
        for p_idx, p in enumerate(shape.text_frame.paragraphs):
            print(f'  P{p_idx}:')
            for r_idx, r in enumerate(p.runs):
                print(f'    R{r_idx}: {repr(r.text)}')
