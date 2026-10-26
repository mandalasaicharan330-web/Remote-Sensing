"""
generate_presentation.py
========================
Builds a professional, high-clarity 10-slide PowerPoint presentation (.pptx)
for the Applied Remote Sensing Final Project:
"Multi-Label Land-Cover Recognition Using Sentinel-2 Multispectral Classification".
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_pptx="presentation/final_project_presentation.pptx"):
    os.makedirs(os.path.dirname(os.path.abspath(output_pptx)), exist_ok=True)
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6] # blank slide layout

    # Color Palette
    DARK_NAVY = RGBColor(22, 36, 54)     # #162436
    LIGHT_BG = RGBColor(248, 250, 252)   # #F8FAFC
    CARD_BG = RGBColor(255, 255, 255)    # #FFFFFF
    PRIMARY_BLUE = RGBColor(30, 75, 124) # #1E4B7C
    ACCENT_TEAL = RGBColor(26, 127, 120) # #1A7F78
    TEXT_DARK = RGBColor(30, 41, 59)     # #1E293B
    TEXT_MUTED = RGBColor(100, 116, 139) # #64748B
    ACCENT_GOLD = RGBColor(217, 119, 6)  # #D97706
    BORDER_COLOR = RGBColor(226, 232, 240)

    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="APPLIED REMOTE SENSING | FINAL PROJECT"):
        # Header banner
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = ACCENT_TEAL

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = PRIMARY_BLUE

    # -------------------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, DARK_NAVY)

    # Accent decorative bar
    bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.6), Inches(1.0), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_GOLD
    bar.line.fill.background()

    tb1 = slide1.shapes.add_textbox(Inches(1.2), Inches(1.9), Inches(11.0), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "Multi-Label Land-Cover Recognition Using Sentinel-2 Multispectral Classification"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.space_after = Pt(14)

    p2 = tf1.add_paragraph()
    p2.text = "Evaluating the Contribution of Visible, NIR, and SWIR Bands on European Satellite Imagery"
    p2.font.size = Pt(18)
    p2.font.color.rgb = RGBColor(203, 213, 225)
    p2.space_after = Pt(28)

    p3 = tf1.add_paragraph()
    p3.text = "Author: Undergraduate Student in Applied Geophysics  |  Antigravity IDE Assistant\n" \
              "Dataset: BigEarthNet v2.0 (Sentinel-2 Level-2A BOA Surface Reflectance)\n" \
              "Date: October 2026"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(148, 163, 184)

    # -------------------------------------------------------------------------
    # SLIDE 2: Background & Motivation
    # -------------------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, LIGHT_BG)
    add_header(slide2, "Background & Motivation: The Mixed Landscape Challenge")

    # Left Column: Problem
    tb2_l = slide2.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.0))
    tf2_l = tb2_l.text_frame
    tf2_l.word_wrap = True

    p = tf2_l.paragraphs[0]
    p.text = "The Single-Label Classification Dilemma"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    p.space_after = Pt(10)

    points_l = [
        ("Landscapes are inherently mixed:", " A single 1.2 km × 1.2 km satellite patch contains pastures bordering woodlots, rivers cutting through fields, and village houses."),
        ("Single-label forces artificial exclusivity:", " Assigning only 'Forest' completely discards secondary classes like pastures or streams."),
        ("Boundary noise & ambiguity:", " Ecotones and mixed pixels create severe labeling conflict in conventional remote sensing.")
    ]
    for bold_prefix, text_body in points_l:
        p = tf2_l.add_paragraph()
        run1 = p.add_run()
        run1.text = "• " + bold_prefix
        run1.font.bold = True
        run1.font.size = Pt(13)
        run1.font.color.rgb = TEXT_DARK
        run2 = p.add_run()
        run2.text = text_body
        run2.font.size = Pt(13)
        run2.font.color.rgb = TEXT_DARK
        p.space_after = Pt(12)

    # Right Column: The Multi-Label Solution
    tb2_r = slide2.shapes.add_textbox(Inches(6.9), Inches(1.7), Inches(5.6), Inches(5.0))
    tf2_r = tb2_r.text_frame
    tf2_r.word_wrap = True

    p = tf2_r.paragraphs[0]
    p.text = "The Multi-Label Paradigm"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TEAL
    p.space_after = Pt(10)

    points_r = [
        ("Independent Bernoulli Formulation:", " Each class is treated independently: y = [y₁, y₂, ..., y₇] where yₖ ∈ {0, 1}."),
        ("82.0% Empirical Multi-Label Prevalence:", " Over 8 out of 10 patches in our dataset contain multiple simultaneous land covers!"),
        ("Realistic Ecological Mapping:", " Accurately detects co-occurring agricultural parcels, woodlands, and built infrastructure without mutual exclusion.")
    ]
    for bold_prefix, text_body in points_r:
        p = tf2_r.add_paragraph()
        run1 = p.add_run()
        run1.text = "• " + bold_prefix
        run1.font.bold = True
        run1.font.size = Pt(13)
        run1.font.color.rgb = TEXT_DARK
        run2 = p.add_run()
        run2.text = text_body
        run2.font.size = Pt(13)
        run2.font.color.rgb = TEXT_DARK
        p.space_after = Pt(12)

    # -------------------------------------------------------------------------
    # SLIDE 3: Objectives & Research Questions
    # -------------------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, LIGHT_BG)
    add_header(slide3, "Research Questions & Project Objectives")

    cards = [
        ("Central Research Question",
         "Can multispectral Sentinel-2 observations (Visible, NIR, SWIR) accurately resolve multiple co-occurring surface types, and how does expanding spectral dimensionality affect multi-label classification accuracy on small curated datasets?",
         PRIMARY_BLUE),
        ("Controlled Band Ablation",
         "Systematically compare 3 spectral configurations under identical architectures and splits:\n"
         "  • M1 (Visible RGB): Bands B02, B03, B04 (3 ch)\n"
         "  • M2 (Visible + NIR): Bands B02, B03, B04, B08 (4 ch)\n"
         "  • M3 (Full Multispectral): B02, B03, B04, B08, B11, B12 (6 ch)",
         ACCENT_TEAL),
        ("Evaluation & Minority Diagnostics",
         "Measure Macro-F1, Micro-F1, and Exact Match subset accuracy. Specifically assess the behavior of minority classes (Wetland, Shrub) and evaluate decision threshold calibration without data leakage.",
         ACCENT_GOLD)
    ]

    for idx, (title, body, color) in enumerate(cards):
        left = Inches(0.8 + idx * 4.0)
        # Background card
        card_shape = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(3.7), Inches(4.8))
        card_shape.fill.solid()
        card_shape.fill.fore_color.rgb = CARD_BG
        card_shape.line.color.rgb = BORDER_COLOR

        tb = slide3.shapes.add_textbox(left + Inches(0.25), Inches(2.0), Inches(3.2), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = color
        p.space_after = Pt(14)

        p2 = tf.add_paragraph()
        p2.text = body
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------------------
    # SLIDE 4: Sentinel-2 Bands & BigEarthNet Dataset
    # -------------------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, LIGHT_BG)
    add_header(slide4, "Sentinel-2 Spectral Bands & BigEarthNet v2.0 Dataset")

    tb4_l = slide4.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(6.0), Inches(5.0))
    tf4_l = tb4_l.text_frame
    tf4_l.word_wrap = True

    p = tf4_l.paragraphs[0]
    p.text = "Diagnostic Spectral Bands Used"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    p.space_after = Pt(8)

    band_rows = [
        ("B02 (Blue - 490 nm, 10m):", " Sensitive to atmospheric haze and water clarity."),
        ("B03 (Green - 560 nm, 10m):", " Captures green chlorophyll peak and vegetation vigor."),
        ("B04 (Red - 665 nm, 10m):", " Maximum chlorophyll-a absorption; defines red edge."),
        ("B08 (NIR - 842 nm, 10m):", " Internal spongy mesophyll leaf scattering; total absorption by open water."),
        ("B11 (SWIR-1 - 1610 nm, 20m):", " Foliar water content, soil moisture, and impervious built-up."),
        ("B12 (SWIR-2 - 2190 nm, 20m):", " Mineral background, dry biomass, and bare soil.")
    ]
    for b_title, b_desc in band_rows:
        p = tf4_l.add_paragraph()
        r1 = p.add_run(); r1.text = "• " + b_title; r1.font.bold = True; r1.font.size = Pt(11); r1.font.color.rgb = TEXT_DARK
        r2 = p.add_run(); r2.text = b_desc; r2.font.size = Pt(11); r2.font.color.rgb = TEXT_DARK
        p.space_after = Pt(4)

    # Right side: dataset stats
    tb4_r = slide4.shapes.add_textbox(Inches(7.2), Inches(1.7), Inches(5.3), Inches(5.0))
    tf4_r = tb4_r.text_frame
    tf4_r.word_wrap = True

    p = tf4_r.paragraphs[0]
    p.text = "Dataset Geography & Partitions"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TEAL
    p.space_after = Pt(8)

    stat_rows = [
        ("Provenance:", " BigEarthNet v2.0 (reBEN), Sentinel-2 Level-2A (Sen2Cor atmospheric correction)."),
        ("Geographic Origin:", " Tile T33UUP in Austrian Central Alps (UTM Zone 33N)."),
        ("Dataset Size:", " 500 genuine patches (1.2 km × 1.2 km, 120×120 pixels)."),
        ("Official Splits:", " 300 Train (60%), 100 Validation (20%), 100 Test (20%)."),
        ("Zero Data Leakage:", " Strictly disjoint partitions; 0 duplicate patch IDs."),
        ("Inspection Patches:", " All 5 Phase 1 inspection patches preserved strictly in Test.")
    ]
    for s_title, s_desc in stat_rows:
        p = tf4_r.add_paragraph()
        r1 = p.add_run(); r1.text = "• " + s_title; r1.font.bold = True; r1.font.size = Pt(11); r1.font.color.rgb = TEXT_DARK
        r2 = p.add_run(); r2.text = s_desc; r2.font.size = Pt(11); r2.font.color.rgb = TEXT_DARK
        p.space_after = Pt(6)

    # -------------------------------------------------------------------------
    # SLIDE 5: Preprocessing & Multispectral Visualization
    # -------------------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, LIGHT_BG)
    add_header(slide5, "Reflectance Preprocessing & Multispectral Views")

    # Add text on left
    tb5 = slide5.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(4.5), Inches(5.0))
    tf5 = tb5.text_frame
    tf5.word_wrap = True

    p = tf5.paragraphs[0]
    p.text = "Harmonization Pipeline"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    p.space_after = Pt(8)

    preproc_points = [
        ("Bilinear Resampling:", " Resampled 20m bands (B11, B12) from 60×60 to 120×120 to align with 10m bands without blocky artifacts."),
        ("Surface Reflectance Scaling:", " ρ = clip(DN / 10000.0, 0.0, 1.0) maps raw DN to physical BOA reflectance [0, 1]."),
        ("Specular Glint Handling:", " Rare reflections > 10,000 DN (<0.0042% of pixels) cleanly bounded without distorting dynamic range."),
        ("True Colour vs False Colour:", " CIR (B8-B4-B3) renders dense biomass in vibrant red, while open water absorbs NIR and appears jet black.")
    ]
    for p_title, p_desc in preproc_points:
        p = tf5.add_paragraph()
        r1 = p.add_run(); r1.text = "• " + p_title; r1.font.bold = True; r1.font.size = Pt(11); r1.font.color.rgb = TEXT_DARK
        r2 = p.add_run(); r2.text = p_desc; r2.font.size = Pt(11); r2.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)

    # Add images on right if they exist
    tc_path = "outputs/figures/phase2_1_true_colour.png"
    fc_path = "outputs/figures/phase2_1_false_colour.png"
    if os.path.exists(tc_path) and os.path.exists(fc_path):
        slide5.shapes.add_picture(tc_path, Inches(5.6), Inches(1.8), width=Inches(3.4))
        slide5.shapes.add_picture(fc_path, Inches(9.3), Inches(1.8), width=Inches(3.4))

    # -------------------------------------------------------------------------
    # SLIDE 6: Seven Target Classes & Class Distribution
    # -------------------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, LIGHT_BG)
    add_header(slide6, "Seven Target Classes & Label Distribution")

    # Table of classes
    rows = 8
    cols = 5
    left = Inches(0.8)
    top = Inches(1.8)
    width = Inches(11.7)
    height = Inches(4.5)

    table_shape = slide6.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    table.columns[0].width = Inches(3.2)
    table.columns[1].width = Inches(2.0)
    table.columns[2].width = Inches(2.0)
    table.columns[3].width = Inches(2.0)
    table.columns[4].width = Inches(2.5)

    headers = ["Target Class Name", "Train (N=300)", "Val (N=100)", "Test (N=100)", "Dataset Prevalence"]
    for c_idx, h in enumerate(headers):
        cell = table.cell(0, c_idx)
        cell.text = h
        cell.fill.solid(); cell.fill.fore_color.rgb = PRIMARY_BLUE
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(11); p.font.bold = True; p.font.color.rgb = RGBColor(255, 255, 255)

    data_rows = [
        ["Forest", "258 (86.0%)", "84 (84.0%)", "95 (95.0%)", "437 (87.4%) - Dominant"],
        ["Agriculture", "142 (47.3%)", "74 (74.0%)", "52 (52.0%)", "268 (53.6%) - Widespread"],
        ["Grassland / Pasture", "209 (69.7%)", "75 (75.0%)", "39 (39.0%)", "323 (64.6%) - Alpine meadows"],
        ["Built-up / Artificial", "130 (43.3%)", "58 (58.0%)", "31 (31.0%)", "219 (43.8%) - Settlements"],
        ["Water", "41 (13.7%)", "28 (28.0%)", "21 (21.0%)", "90 (18.0%) - Lakes & rivers"],
        ["Wetland", "8 (2.7%) [RARE]", "16 (16.0%)", "27 (27.0%)", "51 (10.2%) - Severe imbalance"],
        ["Shrub / Transitional", "59 (19.7%)", "13 (13.0%)", "42 (42.0%)", "114 (22.8%) - Clear-cuts/scrub"]
    ]

    for r_idx, row_data in enumerate(data_rows):
        for c_idx, val in enumerate(row_data):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = val
            cell.fill.solid()
            if r_idx % 2 == 0:
                cell.fill.fore_color.rgb = RGBColor(241, 245, 249)
            else:
                cell.fill.fore_color.rgb = RGBColor(255, 255, 255)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(10)
                p.font.color.rgb = TEXT_DARK
                if "RARE" in val:
                    p.font.bold = True
                    p.font.color.rgb = RGBColor(185, 28, 28)

    # -------------------------------------------------------------------------
    # SLIDE 7: CNN Architecture & Methodology
    # -------------------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, LIGHT_BG)
    add_header(slide7, "CNN Architecture & Multi-Label Training Methodology")

    tb7_l = slide7.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(5.8), Inches(5.0))
    tf7_l = tb7_l.text_frame
    tf7_l.word_wrap = True

    p = tf7_l.paragraphs[0]
    p.text = "MultispectralCNN Architecture"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    p.space_after = Pt(8)

    arch_points = [
        ("Input Flexibility:", " Dynamically accepts C ∈ {3, 4, 6} channels at 120×120 resolution."),
        ("4 Convolutional Blocks:", " [32, 64, 128, 256 filters] with Batch Normalization, ReLU, and Max Pooling (120→60→30→15→7)."),
        ("Spatial Pooling & Dropout:", " AdaptiveAvgPool2d(1, 1) + dual dropout (p=0.4, 0.3) to prevent overfitting on 300 patches."),
        ("Parameter Count:", " Lightweight model (~617k parameters), ideal for small remote sensing datasets."),
        ("Independent Logits:", " Outputs 7 raw unbounded logits, mapped via Sigmoid to [0, 1].")
    ]
    for a_title, a_desc in arch_points:
        p = tf7_l.add_paragraph()
        r1 = p.add_run(); r1.text = "• " + a_title; r1.font.bold = True; r1.font.size = Pt(11); r1.font.color.rgb = TEXT_DARK
        r2 = p.add_run(); r2.text = a_desc; r2.font.size = Pt(11); r2.font.color.rgb = TEXT_DARK
        p.space_after = Pt(6)

    tb7_r = slide7.shapes.add_textbox(Inches(7.0), Inches(1.7), Inches(5.5), Inches(5.0))
    tf7_r = tb7_r.text_frame
    tf7_r.word_wrap = True

    p = tf7_r.paragraphs[0]
    p.text = "Training Protocol & Validation Calibration"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TEAL
    p.space_after = Pt(8)

    train_points = [
        ("Loss Function:", " BCEWithLogitsLoss evaluates 7 independent binary cross-entropy losses simultaneously."),
        ("Optimizer & LR Schedule:", " AdamW (lr=1e-3, weight_decay=1e-4) with CosineAnnealingLR (T_max=35)."),
        ("Validation Early Stopping:", " Monitored validation loss with patience of 8 epochs (saved best weights)."),
        ("Threshold Calibration:", " Tuned per-class thresholds over [0.2, 0.8] strictly on validation data."),
        ("Zero Test Leakage:", " Model weights, normalization, and decision thresholds NEVER touched test labels.")
    ]
    for t_title, t_desc in train_points:
        p = tf7_r.add_paragraph()
        r1 = p.add_run(); r1.text = "• " + t_title; r1.font.bold = True; r1.font.size = Pt(11); r1.font.color.rgb = TEXT_DARK
        r2 = p.add_run(); r2.text = t_desc; r2.font.size = Pt(11); r2.font.color.rgb = TEXT_DARK
        p.space_after = Pt(6)

    # -------------------------------------------------------------------------
    # SLIDE 8: M1 / M2 / M3 Ablation Comparison
    # -------------------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8, LIGHT_BG)
    add_header(slide8, "Experimental Results: M1 vs M2 vs M3 Band Ablation")

    # Add comparison figure on right
    macro_fig = "outputs/figures/phase3_macro_f1_comparison.png"
    if os.path.exists(macro_fig):
        slide8.shapes.add_picture(macro_fig, Inches(6.0), Inches(1.8), width=Inches(6.5))

    # Add text & table on left
    tb8 = slide8.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(5.0), Inches(5.2))
    tf8 = tb8.text_frame
    tf8.word_wrap = True

    p = tf8.paragraphs[0]
    p.text = "Test Evaluation Results (T=0.5)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    p.space_after = Pt(8)

    results_notes = [
        ("M1 (Visible RGB):", " Macro-F1 = 0.5978, Micro-F1 = 0.7403, Exact Match = 12.0%. Baseline optical model."),
        ("M2 (Visible + NIR):", " Macro-F1 = 0.7110 (+11.32 pp), Micro-F1 = 0.8205 (+8.02 pp), Exact Match = 38.0% (+26 pp). WINNER!"),
        ("M3 (Multispectral):", " Macro-F1 = 0.5673, Micro-F1 = 0.7081, Exact Match = 5.0%. Underperformed due to dimensionality & overfitting."),
        ("Validation-Tuned:", " M1 Tuned = 0.7076, M2 Tuned = 0.7020, M3 Tuned = 0.6559. Calibrating thresholds closed M1's gap by lowering cutoffs on minority classes.")
    ]
    for r_title, r_desc in results_notes:
        p = tf8.add_paragraph()
        r1 = p.add_run(); r1.text = "• " + r_title; r1.font.bold = True; r1.font.size = Pt(11); r1.font.color.rgb = TEXT_DARK
        r2 = p.add_run(); r2.text = r_desc; r2.font.size = Pt(11); r2.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)

    # -------------------------------------------------------------------------
    # SLIDE 9: Per-Class Performance, Errors & Wetland Limitation
    # -------------------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9, LIGHT_BG)
    add_header(slide9, "Per-Class Performance, Error Analysis & Wetland Failure")

    # Add per-class figure on right
    pc_fig = "outputs/figures/phase3_per_class_f1_comparison.png"
    if os.path.exists(pc_fig):
        slide9.shapes.add_picture(pc_fig, Inches(6.0), Inches(1.8), width=Inches(6.5))

    tb9 = slide9.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(5.0), Inches(5.2))
    tf9 = tb9.text_frame
    tf9.word_wrap = True

    p = tf9.paragraphs[0]
    p.text = "Key Per-Class Findings"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    p.space_after = Pt(8)

    class_insights = [
        ("Forest & Agriculture:", " Very stable across all models (Forest F1 > 0.96; Agriculture F1 ~0.77-0.84)."),
        ("Shrub Breakthrough in M2:", " F1 surged from 0.3529 (M1) to 0.9114 (M2)! NIR mesophyll scattering separates scrub from mature tree canopy."),
        ("Water Recognition in M2:", " F1 rose from 0.7429 (M1) to 0.8947 (M2) with 100% precision due to NIR water absorption."),
        ("M3 Grassland Over-Prediction:", " M3 produced 57 False Positives on Grassland (precision collapsed to 40%), explaining M3's lower score."),
        ("The Wetland Failure (0.0000 F1):", " All 3 models scored 0 TP on Wetland at T=0.5. With only 8 positive training patches (2.7%), unweighted BCE suppressed predictions.")
    ]
    for c_title, c_desc in class_insights:
        p = tf9.add_paragraph()
        r1 = p.add_run(); r1.text = "• " + c_title; r1.font.bold = True; r1.font.size = Pt(11); r1.font.color.rgb = TEXT_DARK
        r2 = p.add_run(); r2.text = c_desc; r2.font.size = Pt(11); r2.font.color.rgb = TEXT_DARK
        p.space_after = Pt(6)

    # -------------------------------------------------------------------------
    # SLIDE 10: Conclusion, Limitations & Future Work
    # -------------------------------------------------------------------------
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10, LIGHT_BG)
    add_header(slide10, "Conclusions, Limitations & Future Directions")

    # 3 Summary columns
    c_data = [
        ("Key Scientific Conclusions",
         "• Multi-label recognition is vital: 82% of European landscape patches have multiple classes.\n\n"
         "• NIR is the most valuable spectral addition: M2 (Visible+NIR) gained +11.32 pp Macro-F1 and tripled Exact Match (38%).\n\n"
         "• SWIR overfit on small samples: 6 channels without pre-training increased variance and degraded performance.",
         PRIMARY_BLUE),
        ("Project Limitations",
         "• Single-Tile Scope: All 500 patches originate from Tile T33UUP in the Austrian Alps. Findings do not generalize globally.\n\n"
         "• Severe Minority Imbalance: Wetland had only 8 training examples, causing zero-recall.\n\n"
         "• 20m Spatial Interpolation: SWIR bands were bilinearly resampled, lacking native 10m high-frequency detail.",
         ACCENT_GOLD),
        ("Future Research Directions",
         "• Positive-Class Weighting / Focal Loss: Compensate for minority classes like Wetland.\n\n"
         "• Transfer Learning: Pre-train on the full continental BigEarthNet archive before fine-tuning.\n\n"
         "• Geographic Expansion: Evaluate generalization across Mediterranean, Scandinavian, and arid biomes.",
         ACCENT_TEAL)
    ]

    for idx, (title, body, color) in enumerate(c_data):
        left = Inches(0.8 + idx * 4.0)
        card_shape = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(3.7), Inches(5.0))
        card_shape.fill.solid()
        card_shape.fill.fore_color.rgb = CARD_BG
        card_shape.line.color.rgb = BORDER_COLOR

        tb = slide10.shapes.add_textbox(left + Inches(0.2), Inches(2.0), Inches(3.3), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = color
        p.space_after = Pt(12)

        p2 = tf.add_paragraph()
        p2.text = body
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_DARK

    prs.save(output_pptx)
    print(f"Presentation successfully saved to: {output_pptx}")
    return len(prs.slides)

if __name__ == "__main__":
    count = create_deck()
    print(f"Total slides generated: {count}")
