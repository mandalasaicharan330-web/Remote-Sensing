"""
verify_corrected_pptx.py
========================
Verifies the corrections on final_project_presentation_corrected.pptx
"""

import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

pptx_path = "final_project_presentation_corrected.pptx"
prs = Presentation(pptx_path)

print("=" * 80)
print(f"VERIFYING: {pptx_path}")
print("=" * 80)

# Check 1: Slide count
print(f"Slide count: {len(prs.slides)} (Expected: 10)")
assert len(prs.slides) == 10, "Expected exactly 10 slides!"

# Check 2: Slide 1 Author details
s1 = prs.slides[0]
s1_text = ""
for shape in s1.shapes:
    if shape.has_text_frame:
        s1_text += " " + shape.text_frame.text

print("\n--- Slide 1 Text Check ---")
assert "Sai Charan (Mandala Sai Charan)" in s1_text, "Missing author name!"
assert "24JE1093" in s1_text, "Missing admission number!"
assert "Applied Remote Sensing Practical" in s1_text, "Missing course!"
assert "Applied Geophysics, IIT (ISM) Dhanbad" in s1_text, "Missing department/institution!"
assert "Antigravity" not in s1_text, "Found banned 'Antigravity' on Slide 1!"
print("[PASSED] Slide 1 contains correct author details and zero 'Antigravity' mentions.")
print("  Extracted:\n ", s1_text.strip().replace("\n", "\n  "))

# Check 3: Slide 2 Dataset Statistic
s2 = prs.slides[1]
s2_text = ""
for shape in s2.shapes:
    if shape.has_text_frame:
        s2_text += " " + shape.text_frame.text

print("\n--- Slide 2 Text Check ---")
assert "94.4% of patches in our curated dataset contain at least two land-cover classes" in s2_text, "Missing 94.4% wording on Slide 2!"
assert "82" not in s2_text, "Found forbidden '82' on Slide 2!"
print("[PASSED] Slide 2 correctly uses 94.4% curated dataset statistic.")

# Check 4: Slide 8 M3 Performance Explanation & Metrics Preservation
s8 = prs.slides[7]
s8_text = ""
for shape in s8.shapes:
    if shape.has_text_frame:
        s8_text += " " + shape.text_frame.text

print("\n--- Slide 8 Text Check ---")
target_s8_phrase = "The six-band model (M3) underperformed compared with the RGB+NIR model (M2) on the observed test metrics. Possible explanations include the small training sample, spectral collinearity, resampling effects, and optimization instability. These explanations were not independently tested."
assert target_s8_phrase in s8_text, "Missing required M3 cautious explanation on Slide 8!"
assert "Underperformed due to dimensionality & overfitting" not in s8_text, "Found old unverified M3 claim on Slide 8!"
# Verify metrics preserved
assert "0.5978" in s8_text, "Missing M1 Macro-F1!"
assert "0.7110" in s8_text, "Missing M2 Macro-F1!"
assert "0.5673" in s8_text, "Missing M3 Macro-F1!"
assert "0.8205" in s8_text, "Missing M2 Micro-F1!"
assert "38.0%" in s8_text, "Missing M2 Exact Match!"
print("[PASSED] Slide 8 contains required cautious M3 explanation and all preserved verified metrics.")

# Check 5: Slide 10 Statistic & Cautious M3 Conclusion
s10 = prs.slides[9]
s10_text = ""
for shape in s10.shapes:
    if shape.has_text_frame:
        s10_text += " " + shape.text_frame.text

print("\n--- Slide 10 Text Check ---")
assert "94.4% of patches in our curated dataset contain at least two land-cover classes" in s10_text, "Missing 94.4% wording on Slide 10!"
assert "82%" not in s10_text, "Found forbidden '82%' on Slide 10!"
assert "82" not in s10_text, "Found forbidden '82' on Slide 10!"
target_s10_phrase = "The six-band model (M3) performed worse than the RGB+NIR model (M2) on the observed test metrics. Small sample size, spectral collinearity, resampling effects, and optimization instability are possible explanations, but they were not independently tested."
assert target_s10_phrase in s10_text, "Missing required cautious M3 conclusion on Slide 10!"
assert "SWIR overfit on small samples" not in s10_text, "Found old unverified claim on Slide 10!"
print("[PASSED] Slide 10 contains 94.4% curated statistic and cautious M3 conclusion.")

print("\n" + "=" * 80)
print("ALL VERIFICATION CHECKS PASSED WITH ZERO DEFECTS.")
print("=" * 80)
